"""Hope Vale (nbn_accelerate) / Burringurrah (rcp) asset-register -> device-inventory.csv extraction.

One-off, NOT general-purpose. Site asset registers have no standard template (see
references/03_asset-register-conventions.md) -- each site's sheet layout was hand-read and hand-mapped
below. Adapting this to a new site means re-reading that site's actual columns, not assuming they match
either of these two. Kept here (rather than discarded) as a worked example and a reusable set of
guardrails: is_ip()/sanip() reject cross-reference placeholders ("As above") and validate IPv4 format
before a link-derived field becomes a device row; the row() helper whitespace-normalizes every extracted
string (a source MAC cell had an embedded newline that corrupted a CSV row before this was added).

Moved here from cambium-swap's session scratchpad 2026-09-17, per operator instruction, once
skill-cambium existed as this knowledge's canonical home. Re-run: `python3 extract-asset-register.py`
(stdlib + openpyxl; writes the two per-site extract CSVs, then folds them into device-inventory.csv by
appending -- see references/04_device-inventory-schema.md for why a re-extraction of an already-loaded
site must remove that site's old rows first).
"""

import csv, re
import openpyxl

INV_PATH = "/Volumes/Data/_ai/_project/project_stuff/apn/cambium-swap/inventory/device-inventory.csv"
HV_OUT = "/Volumes/Data/_ai/_project/project_stuff/apn/cambium-swap/inventory/asset-register/nbn_accelerate/hope-vale-extract.csv"
BUR_OUT = "/Volumes/Data/_ai/_project/project_stuff/apn/cambium-swap/inventory/asset-register/rcp/burringurrah-extract.csv"
EXPORT_DATE = "2026-09-17"

HEADER = ["serial_msn","mac_address","model","product_family","hardware_revision",
    "current_firmware","previous_known_good_firmware","management_ip","site","smc_flavour",
    "network","tower_site_hierarchy","device_name","device_role","topology_role",
    "wlan_profile","ap_group","configuration_template","cloud_sync_status","cnmaestro_server",
    "licence_tier","licence_state","licence_expiry","local_credentials_ref","last_seen",
    "status","spare_mapping","export_date"]

CRED = {
    "Enterprise Wi-Fi": "<secret:keepassxc:cambium-devices/enterprise-wifi>",
    "ePMP AP": "<secret:keepassxc:cambium-devices/epmp-ap>",
    "ePMP SM": "<secret:keepassxc:cambium-devices/epmp-sm>",
    "cnWave 60 GHz": "<secret:keepassxc:cambium-devices/cnwave-60ghz>",
    "cnPilot R-series": "<secret:keepassxc:cambium-devices/cnpilot-r-series>",
}

U = "UNKNOWN"
rows = []

def row(**kw):
    r = {h: U for h in HEADER}
    r["hardware_revision"] = U
    r["current_firmware"] = U
    r["previous_known_good_firmware"] = U
    r["network"] = U
    r["wlan_profile"] = U
    r["ap_group"] = U
    r["configuration_template"] = U
    r["cloud_sync_status"] = U
    r["cnmaestro_server"] = U
    r["licence_tier"] = U
    r["licence_state"] = U
    r["licence_expiry"] = U
    r["last_seen"] = U
    r["spare_mapping"] = U
    r["export_date"] = EXPORT_DATE
    r["status"] = "Active"
    for k, v in kw.items():
        if isinstance(v, str):
            kw[k] = " ".join(v.split())
    r.update(kw)
    fam = r.get("product_family")
    if r.get("local_credentials_ref") == U and fam in CRED:
        r["local_credentials_ref"] = CRED[fam]
    rows.append(r)

IPV4_RE = re.compile(r"^\d{1,3}(\.\d{1,3}){3}$")

def is_ip(v):
    return bool(v) and bool(IPV4_RE.match(str(v).strip()))

def sanip(v):
    return v if is_ip(v) else U

def norm_f300(text):
    if not text:
        return None
    t = str(text)
    if "300-25" in t:
        return "Force 300-25"
    if "300-16" in t:
        return "Force 300-16"
    return None

# ---------------------------------------------------------------------------
# HOPE VALE (nbn_accelerate)
# ---------------------------------------------------------------------------
hv_path = "/Volumes/Data/_ai/_project/project_stuff/apn/cambium-swap/inventory/asset-register/nbn_accelerate/Hope Vale Asset Register_V1.2.xlsx"
wb = openpyxl.load_workbook(hv_path, data_only=True, read_only=True)
ws = wb["Infrastructure & AP's"]
r_ = list(ws.iter_rows(values_only=True))

SITE = "hope-vale"
FLAVOUR = "nbn_accelerate"

# XV2 AP table: header idx35, data 36-68
for i in range(36, 69):
    rr = r_[i]
    tower, name, model, ip = rr[0], rr[1], rr[2], rr[3]
    backhaul_type, backhaul_name, backhaul_ip = rr[4], rr[5], rr[6]
    if name:
        row(model=model, product_family="Enterprise Wi-Fi", management_ip=sanip(ip), site=SITE,
            smc_flavour=FLAVOUR, tower_site_hierarchy=tower, device_name=name,
            device_role="Wi-Fi AP", topology_role="Enterprise Wi-Fi AP")
    f300_model = norm_f300(backhaul_type)
    if f300_model and backhaul_name and is_ip(backhaul_ip):
        row(model=f300_model, product_family="ePMP SM", management_ip=backhaul_ip, site=SITE,
            smc_flavour=FLAVOUR, tower_site_hierarchy=tower, device_name=backhaul_name,
            device_role="ePMP SM (backhaul)", topology_role="P2P backhaul (paired with XV2 AP)")

# ePMP 3000L table: header idx72, data 73-78
for i in range(73, 79):
    rr = r_[i]
    tower, model, antenna, name, ip, bridge_ssid, freq, backhaul_type = rr[0:8]
    if not name:
        continue
    status = "Spare" if str(tower).strip().lower() == "spare" else "Active"
    row(model="ePMP 3000L", product_family="ePMP AP", management_ip=sanip(ip), site=SITE,
        smc_flavour=FLAVOUR, tower_site_hierarchy=tower, device_name=name,
        device_role="ePMP AP (P2MP)", topology_role=antenna, status=status)

# 5GHz ePMP P2P Links: header idx81, data row 82 (AP + SM per row)
for i in range(82, 86):
    rr = r_[i]
    if not any(rr):
        continue
    ap_tower, ap_model, ap_name, ap_ip, bridge_ssid, freq, chwidth, sm_tower, sm_model, sm_name, sm_ip = rr[0:11]
    status = "Spare" if str(ap_tower).strip().lower() == "spare" else "Active"
    if ap_name:
        row(model=norm_f300(ap_model) or ap_model, product_family="ePMP SM", management_ip=sanip(ap_ip),
            site=SITE, smc_flavour=FLAVOUR, tower_site_hierarchy=ap_tower, device_name=ap_name,
            device_role="ePMP AP (P2P master)", topology_role="P2P AP", status=status)
    if sm_name:
        row(model=norm_f300(sm_model) or sm_model, product_family="ePMP SM", management_ip=sanip(sm_ip),
            site=SITE, smc_flavour=FLAVOUR, tower_site_hierarchy=sm_tower, device_name=sm_name,
            device_role="ePMP SM (P2P slave)", topology_role="P2P SM", status=status)

# 60GHz P2(M)P Links: header idx87, data 88-96 (DN once at 88, CN per row)
for i in range(88, 97):
    rr = r_[i]
    tower, dn_site, dn_name, dn_ip, dn_model, cn_loc, cn_name, cn_ip, cn_model, cn_sitename, sector, cn_mac = rr[0:12]
    if dn_name:
        row(model=dn_model, product_family="cnWave 60 GHz", management_ip=sanip(dn_ip), site=SITE,
            smc_flavour=FLAVOUR, tower_site_hierarchy=tower or "Tower 1", device_name=dn_name,
            device_role="cnWave DN (P2MP hub)", topology_role="DN")
    if cn_name:
        row(model=cn_model, product_family="cnWave 60 GHz", management_ip=sanip(cn_ip),
            mac_address=cn_mac or U, site=SITE, smc_flavour=FLAVOUR,
            tower_site_hierarchy=cn_loc, device_name=cn_name,
            device_role="cnWave CN (P2MP client)", topology_role="CN")

# 60GHz Spare DN's: header idx100, data 101
rr = r_[101]
tower, dn_site, dn_name, dn_ip, dn_model = rr[0:5]
if dn_name:
    row(model=dn_model, product_family="cnWave 60 GHz", management_ip=sanip(dn_ip), site=SITE,
        smc_flavour=FLAVOUR, tower_site_hierarchy=tower, device_name=dn_name,
        device_role="cnWave DN (spare)", topology_role="DN", status="Spare")

hv_count = len(rows)
print(f"Hope Vale rows: {hv_count}")

# ---------------------------------------------------------------------------
# BURRINGURRAH (rcp)
# ---------------------------------------------------------------------------
bur_path = "/Volumes/Data/_ai/_project/project_stuff/apn/cambium-swap/inventory/asset-register/rcp/Burringurrah Asset Register.xlsx"
wb2 = openpyxl.load_workbook(bur_path, data_only=True, read_only=True)

SITE = "burringurrah"
FLAVOUR = "rcp"

ws2 = wb2["Tower + AP"]
r2 = list(ws2.iter_rows(values_only=True))

# Combined 3000L omni + Force300 SM block: header idx0, data 1-9
for i in range(1, 10):
    rr = r2[i]
    lat, lon, tid, loc, name, ip, serial, mac = rr[0:8]
    bridge_ssid, freq, bridge_mac, power, mast = rr[10:15]
    if not name:
        continue
    latest_name = str(name).split(">>")[-1].strip()
    f300_model = norm_f300(latest_name)
    status = "Spare" if "spare" in str(name).lower() else "Active"
    if f300_model:
        row(mac_address=mac or U, model=f300_model, product_family="ePMP SM", management_ip=sanip(ip),
            site=SITE, smc_flavour=FLAVOUR, tower_site_hierarchy=tid or loc, device_name=latest_name,
            device_role="ePMP SM (backhaul)", topology_role="P2P backhaul", status=status)
    else:
        row(mac_address=mac or U, model="ePMP 3000L", product_family="ePMP AP", management_ip=sanip(ip),
            site=SITE, smc_flavour=FLAVOUR, tower_site_hierarchy=tid or loc, device_name=latest_name,
            device_role="ePMP AP (P2MP)", topology_role="Omni/Sector", status=status)

# XV2 AP table: header idx11, data 13-20
for i in range(13, 21):
    rr = r2[i]
    lat, lon, apid, loc, name, ip = rr[0:6]
    ap_serial, ap_mac = rr[7], rr[8]
    if not name:
        continue
    status = "Spare" if str(apid).strip().upper() == "SPARE" else "Active"
    row(mac_address=ap_mac or U, model="XV2", product_family="Enterprise Wi-Fi", management_ip=sanip(ip),
        site=SITE, smc_flavour=FLAVOUR, tower_site_hierarchy=apid, device_name=name,
        device_role="Wi-Fi AP", topology_role="Enterprise Wi-Fi AP (variant XV2-2T0 vs XV2-22H unconfirmed)",
        status=status)

# Switches sheet skipped (non-Cambium, out of scope)

bur_towerap_count = len(rows) - hv_count

# Internals sheet: residential R195P CPE installs
ws3 = wb2["Internals"]
r3 = list(ws3.iter_rows(values_only=True))
# header at row0: idx2 Street Number, idx3 Street Address, idx13 R195 EXT, idx27 R195 MAC Address
r195_count = 0
for i in range(1, len(r3)):
    rr = r3[i]
    if len(rr) <= 27:
        continue
    street_no, street_addr = rr[2], rr[3]
    r195_ext = rr[13]
    r195_mac = rr[27]
    if r195_ext is None:
        continue
    hierarchy = f"{street_no} {street_addr}".strip() if street_addr else f"EXT {r195_ext}"
    # Operator-stated derivation (2026-09-17, not in the register): EXT10XX -> 10.255.10.XX
    ext_str = str(r195_ext).strip()
    ext_digits = re.match(r"^10(\d{2})", ext_str)
    r195_ip = f"10.255.10.{int(ext_digits.group(1))}" if ext_digits else U
    row(mac_address=r195_mac or U, model="R195P", product_family="cnPilot R-series",
        management_ip=r195_ip, site=SITE, smc_flavour=FLAVOUR, tower_site_hierarchy=hierarchy,
        device_name=f"R195P EXT{r195_ext}", device_role="CPE (R195P)",
        topology_role="Residential CPE")
    r195_count += 1

print(f"Burringurrah Tower+AP rows: {bur_towerap_count}")
print(f"Burringurrah R195P CPE rows: {r195_count}")
print(f"TOTAL new rows: {len(rows)}")

hv_rows = rows[:hv_count]
bur_rows = rows[hv_count:]

# ---------------------------------------------------------------------------
# Per-site normalized extracts (reviewable checkpoints, re-runnable)
# ---------------------------------------------------------------------------
for path, site_rows in [(HV_OUT, hv_rows), (BUR_OUT, bur_rows)]:
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=HEADER, quoting=csv.QUOTE_MINIMAL)
        w.writeheader()
        for r in site_rows:
            w.writerow(r)
    print(f"Wrote {path} ({len(site_rows)} rows)")

# ---------------------------------------------------------------------------
# Fold both extracts into device-inventory.csv (append, header stays as-is)
# ---------------------------------------------------------------------------
with open(INV_PATH, "a", newline="") as f:
    w = csv.DictWriter(f, fieldnames=HEADER, quoting=csv.QUOTE_MINIMAL)
    for r in rows:
        w.writerow(r)

print("Folded into device-inventory.csv.")
