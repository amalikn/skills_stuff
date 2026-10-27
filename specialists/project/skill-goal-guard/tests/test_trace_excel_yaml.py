"""Tests for milestone trace (goal_guard_lib.trace) and the two library seams' upgrade tests (yaml_io for PyYAML, excel for openpyxl).

Run: just test
Safety: files only, in temporary folders.
"""

import datetime as dt
import tempfile
import unittest
from pathlib import Path

from helpers import Workspace, lib, tasks
from openpyxl import load_workbook

trace = lib("trace")
excel = lib("excel")
yaml_io = lib("yaml_io")


class Trace(unittest.TestCase):
  """A verified link, and each way it becomes REVIEW."""

  def run_trace(self, ws: Workspace) -> dict:
    """Trace milestone M1 of the spoke.

    Args:
      ws: the workspace.

    Returns:
      The trace result.
    """
    proj = ws.project()
    return trace.trace_milestone(proj, proj.milestones[0], tasks(), 5)

  @staticmethod
  def failed(result: dict) -> list[str]:
    """List the failed checks.

    Args:
      result: a trace result.

    Returns:
      Names of failed checks.
    """
    return [c["check"] for c in result["checks"] if not c["ok"]]

  def test_fixture_verifies(self):
    """The example link passes every check."""
    with Workspace() as ws:
      self.assertEqual(self.run_trace(ws)["result"], "verified")

  def test_changed_plan_reviews(self):
    """An edited hub plan no longer matches the pinned revision."""
    with Workspace() as ws:
      plan = ws.root / "apn/demo-hub/plan.md"
      plan.write_text(plan.read_text() + "\nchanged\n")
      result = self.run_trace(ws)
    self.assertEqual((result["result"], self.failed(result)), ("REVIEW", ["plan_baseline"]))

  def test_unknown_target_and_task_review(self):
    """A target ID missing from the map and a task missing from Vikunja both fail."""
    with Workspace() as ws:
      proj = ws.project()
      link = dict(proj.milestones[0], target_id="T-NOPE", vikunja_task_ids=[101, 999])
      result = trace.trace_milestone(proj, link, tasks(), 5)
    self.assertTrue({"target_id", "vikunja_tasks"} <= set(self.failed(result)))


class Seams(unittest.TestCase):
  """Upgrade tests: fail if a PyYAML or openpyxl upgrade changes the behaviour relied on."""

  def test_yaml_safe_load_refuses_python_tags(self):
    """safe_load rejects a python tag instead of constructing an object."""
    with tempfile.TemporaryDirectory() as tmp:
      path = Path(tmp) / "x.yaml"
      path.write_text("!!python/object/apply:builtins.print ['x']\n")
      with self.assertRaises(lib("errors").RecordError):
        yaml_io.load_file(path)

  def test_yaml_dump_keeps_order(self):
    """dump_text keeps insertion order."""
    self.assertTrue(yaml_io.dump_text({"b": 1, "a": "2026-12-01"}).startswith("b: 1"))

  def test_workbook_columns_and_privacy(self):
    """The sample has the default columns, no internal target unless asked, and provenance on the About sheet."""
    with Workspace() as ws, tempfile.TemporaryDirectory() as tmp:
      proj = ws.project()
      out = Path(tmp) / "s.xlsx"
      excel.write(out, proj, excel.rows(proj, dt.date(2026, 12, 5), False), "2026-12-05T09:00:00+11:00")
      book = load_workbook(out)
      header = [c.value for c in book["Milestones"][1]]
      about = {r[0].value: r[1].value for r in book["About"].iter_rows()}
      self.assertEqual(header, list(excel.COLUMNS))
      self.assertEqual(book["Milestones"]["F2"].value, "deadline at risk")
      self.assertEqual(about["mvp_revision"], "DEMO-MVP-001@1")
      self.assertEqual(excel.rows(proj, dt.date(2026, 12, 5), True)[0]["internal_target"], "2026-12-01")


if __name__ == "__main__":
  unittest.main()
